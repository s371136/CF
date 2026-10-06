# -*- coding: utf-8 -*-
"""
CREA INPUT HORNSUND 8760 h - VERSIONE CON FILE LOAD SEPARATO

Unisce:
- Load [kW] da un file Excel selezionato dall'utente
- CF WIND dell'anno selezionato
- CF PV dello stesso anno
- Timestamp ricreati usando l'anno selezionato per PV/Wind

Output:
- un solo foglio: "Input data"
- 8760 righe orarie
- layout iniziale come il file Hornsund mostrato (titolo, nota, location)

Nota importante:
- il file LOAD serve per importare il profilo di carico di 8760 ore
- i timestamp dell'output vengono sempre ricreati con l'anno selezionato
  per i profili PV/Wind; per gli anni bisestili viene eliminato il 29 febbraio
"""

from __future__ import annotations

from pathlib import Path
import sys
import traceback

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# =============================================================================
# SETTINGS
# =============================================================================

INPUT_SHEET = "Input data"
START_ROW = 6
N_HOURS = 8760
END_ROW = START_ROW + N_HOURS - 1

LATITUDE = 77.00213
LONGITUDE = 15.53979

TITLE_FILL = "17365D"      # blu scuro simile al file mostrato
HEADER_FILL = "D9EAF7"     # azzurro chiaro per intestazioni
WHITE = "FFFFFF"
BLACK = "000000"


# =============================================================================
# TKINTER
# =============================================================================

def _new_hidden_root():
    import tkinter as tk

    root = tk.Tk()
    root.withdraw()

    try:
        root.attributes("-topmost", True)
    except Exception:
        pass

    return root


def select_excel_file(title: str) -> Path:
    from tkinter import filedialog

    root = _new_hidden_root()

    try:
        selected = filedialog.askopenfilename(
            parent=root,
            title=title,
            filetypes=[
                ("File Excel", "*.xlsx"),
                ("Tutti i file", "*.*"),
            ],
        )
    finally:
        root.destroy()

    if not selected:
        raise SystemExit("Selezione file annullata.")

    return Path(selected)


def select_output_file(default_directory: Path, default_filename: str) -> Path:
    from tkinter import filedialog

    root = _new_hidden_root()

    try:
        selected = filedialog.asksaveasfilename(
            parent=root,
            title="Salva il nuovo input Hornsund",
            initialdir=str(default_directory),
            initialfile=default_filename,
            defaultextension=".xlsx",
            filetypes=[
                ("File Excel", "*.xlsx"),
                ("Tutti i file", "*.*"),
            ],
        )
    finally:
        root.destroy()

    if not selected:
        raise SystemExit("Salvataggio annullato.")

    output = Path(selected)

    if output.suffix.lower() != ".xlsx":
        output = output.with_suffix(".xlsx")

    return output


def show_info(title: str, message: str) -> None:
    from tkinter import messagebox

    root = _new_hidden_root()
    try:
        messagebox.showinfo(title, message, parent=root)
    finally:
        root.destroy()


def show_error(title: str, message: str) -> None:
    from tkinter import messagebox

    root = _new_hidden_root()
    try:
        messagebox.showerror(title, message, parent=root)
    finally:
        root.destroy()


def choose_year_dialog(years: list[int]) -> int:
    """Sceglie l'anno comune tra i file PV e Wind."""
    if not years:
        raise ValueError("Non ci sono anni comuni tra PV e wind.")

    if len(years) == 1:
        return int(years[0])

    import tkinter as tk
    from tkinter import ttk

    result = {"year": None}

    dialog = tk.Tk()
    dialog.title("Anno da usare")
    dialog.resizable(False, False)

    try:
        dialog.attributes("-topmost", True)
    except Exception:
        pass

    frame = ttk.Frame(dialog, padding=16)
    frame.grid(row=0, column=0)

    ttk.Label(
        frame,
        text="Seleziona l'anno dei profili PV e Wind:",
    ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))

    labels = [str(year) for year in years]
    variable = tk.StringVar(value=labels[0])

    combo = ttk.Combobox(
        frame,
        textvariable=variable,
        values=labels,
        state="readonly",
        width=15,
    )
    combo.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 12))

    def accept():
        result["year"] = int(variable.get())
        dialog.destroy()

    def cancel():
        dialog.destroy()

    ttk.Button(frame, text="OK", command=accept).grid(row=2, column=0, padx=(0, 5))
    ttk.Button(frame, text="Annulla", command=cancel).grid(row=2, column=1, padx=(5, 0))

    dialog.protocol("WM_DELETE_WINDOW", cancel)
    dialog.bind("<Return>", lambda _event: accept())
    dialog.bind("<Escape>", lambda _event: cancel())

    dialog.update_idletasks()
    x = (dialog.winfo_screenwidth() - dialog.winfo_reqwidth()) // 2
    y = (dialog.winfo_screenheight() - dialog.winfo_reqheight()) // 2
    dialog.geometry(f"+{x}+{y}")

    dialog.lift()
    dialog.focus_force()
    combo.focus_set()
    dialog.mainloop()

    if result["year"] is None:
        raise SystemExit("Operazione annullata.")

    return int(result["year"])


# =============================================================================
# TIME INDEX / NORMALIZZAZIONE
# =============================================================================

def expected_index_8760(year: int) -> pd.DatetimeIndex:
    """Crea l'indice orario dell'anno e rimuove il 29 febbraio se presente."""
    index = pd.date_range(
        start=f"{year}-01-01 00:00:00",
        end=f"{year}-12-31 23:00:00",
        freq="h",
    )

    if len(index) == 8784:
        index = index[~((index.month == 2) & (index.day == 29))]

    if len(index) != N_HOURS:
        raise ValueError(
            f"L'anno {year} produce {len(index)} ore invece di {N_HOURS}."
        )

    return index


def _normalise_timestamp(series: pd.Series) -> pd.Series:
    timestamps = pd.to_datetime(series, errors="coerce")

    try:
        if timestamps.dt.tz is not None:
            timestamps = timestamps.dt.tz_localize(None)
    except AttributeError:
        pass

    return timestamps.dt.floor("h")


# =============================================================================
# LOAD - FILE SEPARATO CON TIMESTAMP + LOAD
# =============================================================================

def _find_column(columns, candidates: list[str]):
    """Trova una colonna in modo tollerante rispetto a maiuscole/spazi/unità."""
    normalized = {
        str(col).strip().lower().replace(" ", ""): col
        for col in columns
    }

    for candidate in candidates:
        key = candidate.strip().lower().replace(" ", "")
        if key in normalized:
            return normalized[key]

    return None


def read_load_file(load_path: Path) -> tuple[pd.DatetimeIndex, np.ndarray]:
    """
    Legge un file Excel con due colonne principali:
    - Timestamp
    - Load / Load [kW]

    Usa il primo foglio del file. I timestamp del LOAD vengono usati solo
    per verificare che il profilo sia orario e continuo; l'anno dell'output
    viene invece imposto in base all'anno selezionato per PV/Wind.
    """
    print("Lettura del file LOAD...", flush=True)

    excel = pd.ExcelFile(load_path, engine="openpyxl")
    if not excel.sheet_names:
        raise ValueError("Il file LOAD non contiene fogli Excel.")

    sheet_name = excel.sheet_names[0]
    df = pd.read_excel(load_path, sheet_name=sheet_name, engine="openpyxl")

    timestamp_col = _find_column(
        df.columns,
        [
            "Timestamp",
            "Date and time",
            "Date and time UTC",
            "Datetime",
            "DateTime",
        ],
    )

    load_col = _find_column(
        df.columns,
        [
            "Load [kW]",
            "Load",
            "Load(kW)",
            "Power [kW]",
            "Demand [kW]",
        ],
    )

    # Fallback utile per i file semplici a 2 colonne creati apposta per questo script.
    if timestamp_col is None or load_col is None:
        if df.shape[1] >= 2:
            print(
                "Intestazioni non riconosciute automaticamente: uso le prime due colonne "
                "come Timestamp e Load.",
                flush=True,
            )
            timestamp_col = df.columns[0]
            load_col = df.columns[1]
        else:
            raise ValueError(
                "Il file LOAD deve contenere almeno due colonne: Timestamp e Load [kW]."
            )

    timestamps = _normalise_timestamp(df[timestamp_col])
    load = pd.to_numeric(df[load_col], errors="coerce")

    data = pd.DataFrame({"timestamp": timestamps, "load": load})
    data = data.dropna(subset=["timestamp", "load"]).copy()

    # Se il file contiene un 29 febbraio, lo eliminiamo per mantenere 8760 ore.
    data = data[
        ~(
            (data["timestamp"].dt.month == 2)
            & (data["timestamp"].dt.day == 29)
        )
    ].copy()

    if len(data) != N_HOURS:
        raise ValueError(
            f"LOAD: trovate {len(data)} ore valide invece di {N_HOURS}."
        )

    if data["timestamp"].duplicated().any():
        n_dup = int(data["timestamp"].duplicated().sum())
        raise ValueError(f"LOAD: trovati {n_dup} timestamp duplicati.")

    if data["load"].isna().any():
        bad = int(data["load"].isna().sum())
        raise ValueError(f"LOAD: {bad} valori non numerici o vuoti.")

    if (data["load"] < 0).any():
        raise ValueError("Il LOAD contiene valori negativi.")

    # Verifica che l'ordine sia temporale e orario.
    data = data.sort_values("timestamp").reset_index(drop=True)
    delta = data["timestamp"].diff().dropna()
    if not (delta == pd.Timedelta(hours=1)).all():
        raise ValueError(
            "I timestamp del LOAD non sono una serie oraria continua da 8760 ore."
        )

    timestamps_out = pd.DatetimeIndex(data["timestamp"])
    load_out = data["load"].to_numpy(dtype=float)

    print(
        f"LOAD letto correttamente: {len(load_out)} ore, "
        f"indice {timestamps_out[0]} -> {timestamps_out[-1]}",
        flush=True,
    )

    return timestamps_out, load_out


# =============================================================================
# PV
# =============================================================================

def get_pv_years(pv_path: Path) -> list[int]:
    excel = pd.ExcelFile(pv_path, engine="openpyxl")
    years = []

    for sheet in excel.sheet_names:
        name = str(sheet).strip()
        if name.lower().startswith("hourly "):
            try:
                years.append(int(name.split()[-1]))
            except ValueError:
                pass

    return sorted(set(years))


def read_pv_year(
    pv_path: Path,
    year: int,
) -> tuple[pd.Series, pd.Series]:
    sheet_name = f"Hourly {year}"

    print(f"Lettura PV {year} dal foglio '{sheet_name}'...", flush=True)

    df = pd.read_excel(
        pv_path,
        sheet_name=sheet_name,
        engine="openpyxl",
    )

    required = ["Date and time UTC", "Capacity factor (kW/kWp)"]
    missing = [column for column in required if column not in df.columns]

    if missing:
        raise ValueError(
            "Nel file PV mancano le colonne: " + ", ".join(missing)
        )

    timestamps = _normalise_timestamp(df["Date and time UTC"])
    cf = pd.to_numeric(df["Capacity factor (kW/kWp)"], errors="coerce")

    if "GHI horizontal (W/m²)" in df.columns:
        ghi = pd.to_numeric(df["GHI horizontal (W/m²)"], errors="coerce")
    else:
        ghi = pd.Series(np.nan, index=df.index)

    data = pd.DataFrame(
        {
            "timestamp": timestamps,
            "cf_pv": cf,
            "ghi": ghi,
        }
    )

    data = data.dropna(subset=["timestamp", "cf_pv"])
    data = data[data["timestamp"].dt.year == year].copy()
    data = data[
        ~(
            (data["timestamp"].dt.month == 2)
            & (data["timestamp"].dt.day == 29)
        )
    ]

    data = (
        data.drop_duplicates(subset="timestamp", keep="first")
        .set_index("timestamp")
        .sort_index()
    )

    expected = expected_index_8760(year)
    data = data.reindex(expected)

    missing_cf = int(data["cf_pv"].isna().sum())
    if missing_cf:
        missing_times = data.index[data["cf_pv"].isna()][:5]
        example = ", ".join(ts.strftime("%Y-%m-%d %H:%M") for ts in missing_times)
        raise ValueError(
            f"PV {year}: mancano {missing_cf} ore. Prime ore mancanti: {example}"
        )

    if (data["cf_pv"] < -1e-9).any() or (data["cf_pv"] > 1.0 + 1e-9).any():
        raise ValueError("Il CF PV contiene valori fuori dal range [0, 1].")

    print(
        f"PV {year}: {len(data)} ore, CF medio = {data['cf_pv'].mean():.4f}",
        flush=True,
    )

    return data["cf_pv"], data["ghi"]


# =============================================================================
# WIND
# =============================================================================

def read_wind_table(wind_path: Path) -> pd.DataFrame:
    print("Lettura del file wind...", flush=True)

    df = pd.read_excel(
        wind_path,
        sheet_name="Hourly profile",
        engine="openpyxl",
    )

    required = ["Date and time", "Capacity factor (-)"]
    missing = [column for column in required if column not in df.columns]

    if missing:
        raise ValueError(
            "Nel file wind mancano le colonne: " + ", ".join(missing)
        )

    df = df.copy()
    df["timestamp"] = _normalise_timestamp(df["Date and time"])
    df["cf_wind"] = pd.to_numeric(df["Capacity factor (-)"], errors="coerce")

    if "Wind speed at turbine height (m/s)" in df.columns:
        df["wind_speed"] = pd.to_numeric(
            df["Wind speed at turbine height (m/s)"], errors="coerce"
        )
    elif "Wind speed at reference height (m/s)" in df.columns:
        df["wind_speed"] = pd.to_numeric(
            df["Wind speed at reference height (m/s)"], errors="coerce"
        )
    else:
        df["wind_speed"] = np.nan

    df = df.dropna(subset=["timestamp", "cf_wind"])

    return df[["timestamp", "cf_wind", "wind_speed"]].copy()


def get_wind_years(wind_df: pd.DataFrame) -> list[int]:
    years = (
        wind_df["timestamp"]
        .dt.year
        .dropna()
        .astype(int)
        .unique()
        .tolist()
    )
    return sorted(years)


def read_wind_year(
    wind_df: pd.DataFrame,
    year: int,
) -> tuple[pd.Series, pd.Series]:
    print(f"Preparazione wind {year}...", flush=True)

    data = wind_df[wind_df["timestamp"].dt.year == year].copy()
    data = data[
        ~(
            (data["timestamp"].dt.month == 2)
            & (data["timestamp"].dt.day == 29)
        )
    ]

    data = (
        data.drop_duplicates(subset="timestamp", keep="first")
        .set_index("timestamp")
        .sort_index()
    )

    expected = expected_index_8760(year)
    data = data.reindex(expected)

    missing_cf = int(data["cf_wind"].isna().sum())
    if missing_cf:
        missing_times = data.index[data["cf_wind"].isna()][:5]
        example = ", ".join(ts.strftime("%Y-%m-%d %H:%M") for ts in missing_times)
        raise ValueError(
            f"WIND {year}: mancano {missing_cf} ore. Prime ore mancanti: {example}"
        )

    if (data["cf_wind"] < -1e-9).any() or (data["cf_wind"] > 1.0 + 1e-9).any():
        raise ValueError("Il CF wind contiene valori fuori dal range [0, 1].")

    print(
        f"WIND {year}: {len(data)} ore, CF medio = {data['cf_wind'].mean():.4f}",
        flush=True,
    )

    return data["cf_wind"], data["wind_speed"]


# =============================================================================
# OUTPUT - SOLO FOGLIO INPUT DATA
# =============================================================================

def write_output(
    output_path: Path,
    profile_year: int,
    load: np.ndarray,
    cf_wind: pd.Series,
    cf_pv: pd.Series,
    wind_speed: pd.Series,
    ghi: pd.Series,
) -> None:
    print("Creazione del file Excel...", flush=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = INPUT_SHEET

    # -------------------------------------------------------------------------
    # Prime celle come nel file mostrato
    # -------------------------------------------------------------------------
    sheet.merge_cells("A1:G1")
    sheet["A1"] = f"Hornsund input - {profile_year} - 8760 h"
    sheet["A1"].font = Font(color=WHITE, bold=True, size=14)
    sheet["A1"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 22

    sheet["A2"] = (
        "Load importato dal file selezionato; CF wind e CF PV sostituiti "
        "con i profili dell'anno indicato."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 62

    # Riga 3 volutamente vuota, come nell'esempio.
    sheet.row_dimensions[3].height = 18

    sheet["A4"] = (
        f"Location: {LATITUDE:.5f}° N, {LONGITUDE:.5f}° E | "
        f"8760 h index: {profile_year}"
    )
    sheet["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[4].height = 42

    headers = [
        "Timestamp",
        "Load [kW]",
        "CF wind [-]",
        "CF PV [-]",
        "Wind speed [m/s]",
        "GHI [W/m²]",
        "Temperature [°C]",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=col_idx, value=header)
        cell.font = Font(bold=True, color=BLACK)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.row_dimensions[5].height = 30

    # Dimensioni colonne simili all'Excel mostrato.
    widths = {
        "A": 23,
        "B": 16,
        "C": 16,
        "D": 16,
        "E": 20,
        "F": 20,
        "G": 20,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    # -------------------------------------------------------------------------
    # Dati 8760 h
    # -------------------------------------------------------------------------
    timestamps = expected_index_8760(profile_year)
    cf_wind_values = cf_wind.to_numpy(dtype=float)
    cf_pv_values = cf_pv.to_numpy(dtype=float)
    wind_values = wind_speed.to_numpy()
    ghi_values = ghi.to_numpy()

    if not (
        len(timestamps)
        == len(load)
        == len(cf_wind_values)
        == len(cf_pv_values)
        == N_HOURS
    ):
        raise ValueError("Le serie da scrivere non hanno tutte 8760 valori.")

    for index in range(N_HOURS):
        row = START_ROW + index

        # Timestamp: usa SEMPRE l'anno selezionato per PV/Wind.
        sheet.cell(row, 1).value = timestamps[index].to_pydatetime()
        sheet.cell(row, 1).number_format = "yyyy-mm-dd hh:mm"

        sheet.cell(row, 2).value = float(load[index])
        sheet.cell(row, 3).value = float(cf_wind_values[index])
        sheet.cell(row, 4).value = float(cf_pv_values[index])

        wind_value = wind_values[index]
        ghi_value = ghi_values[index]

        sheet.cell(row, 5).value = None if pd.isna(wind_value) else float(wind_value)
        sheet.cell(row, 6).value = None if pd.isna(ghi_value) else float(ghi_value)

        # Il nuovo file LOAD contiene solo timestamp + load, quindi temperatura vuota.
        sheet.cell(row, 7).value = None

    # Formati numerici
    for row in range(START_ROW, END_ROW + 1):
        sheet.cell(row, 2).number_format = "0.000"
        sheet.cell(row, 3).number_format = "0.000000"
        sheet.cell(row, 4).number_format = "0.000000"
        sheet.cell(row, 5).number_format = "0.000"
        sheet.cell(row, 6).number_format = "0.000"

    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:G{END_ROW}"

    print("Salvataggio del file Excel...", flush=True)
    workbook.save(output_path)
    workbook.close()
    print("File Excel salvato.", flush=True)


# =============================================================================
# MAIN
# =============================================================================

def main() -> None:
    print("=" * 78, flush=True)
    print("HORNSUND INPUT BUILDER - LOAD FILE + PV + WIND - 8760 h", flush=True)
    print("=" * 78, flush=True)

    # 1. LOAD
    print("1. Seleziona il file LOAD (Timestamp + Load)...", flush=True)
    load_path = select_excel_file(
        "1/3 - Seleziona il file LOAD con Timestamp e Load [kW]"
    )
    print(f"LOAD: {load_path}", flush=True)

    # 2. PV
    print("2. Seleziona il file PV...", flush=True)
    pv_path = select_excel_file(
        "2/3 - Seleziona il file PV capacity factor"
    )
    print(f"PV: {pv_path}", flush=True)

    # 3. WIND
    print("3. Seleziona il file wind...", flush=True)
    wind_path = select_excel_file(
        "3/3 - Seleziona il file Wind capacity factor"
    )
    print(f"Wind: {wind_path}", flush=True)

    # 4. Anni disponibili
    print("4. Lettura degli anni disponibili...", flush=True)
    pv_years = get_pv_years(pv_path)
    wind_df = read_wind_table(wind_path)
    wind_years = get_wind_years(wind_df)

    print(f"Anni PV: {pv_years}", flush=True)
    print(f"Anni wind: {wind_years}", flush=True)

    common_years = sorted(set(pv_years) & set(wind_years))
    if not common_years:
        raise ValueError(
            "Nessun anno comune tra i due file.\n"
            f"PV: {pv_years}\n"
            f"Wind: {wind_years}"
        )

    profile_year = choose_year_dialog(common_years)
    print(f"Anno profili selezionato: {profile_year}", flush=True)

    # 5. LOAD
    print("5. Lettura del load...", flush=True)
    load_timestamps, load = read_load_file(load_path)

    # 6. CF
    print("6. Lettura dei capacity factor...", flush=True)
    cf_pv, ghi = read_pv_year(pv_path, profile_year)
    cf_wind, wind_speed = read_wind_year(wind_df, profile_year)

    # 7. Output
    print("7. Seleziona nome e cartella del file finale...", flush=True)
    output_path = select_output_file(
        load_path.parent,
        f"Hornsund_input_{profile_year}_8760.xlsx",
    )
    print(f"Output: {output_path}", flush=True)

    # 8. Scrittura
    print("8. Creazione del file finale...", flush=True)
    write_output(
        output_path=output_path,
        profile_year=profile_year,
        load=load,
        cf_wind=cf_wind,
        cf_pv=cf_pv,
        wind_speed=wind_speed,
        ghi=ghi,
    )

    annual_load_mwh = float(load.sum() / 1000.0)
    mean_wind_cf = float(cf_wind.mean())
    mean_pv_cf = float(cf_pv.mean())

    print("\n" + "=" * 78, flush=True)
    print("COMPLETATO", flush=True)
    print("=" * 78, flush=True)
    print(f"Anno profili PV/Wind: {profile_year}", flush=True)
    print(f"Anno indice output: {profile_year}", flush=True)
    print(f"Ore: {N_HOURS}", flush=True)
    print(f"Load annuale: {annual_load_mwh:.3f} MWh", flush=True)
    print(f"CF wind medio: {mean_wind_cf:.4f}", flush=True)
    print(f"CF PV medio: {mean_pv_cf:.4f}", flush=True)
    print(f"Creato: {output_path}", flush=True)

    show_info(
        "File Hornsund creato",
        (
            f"Anno profili PV/Wind: {profile_year}\n"
            f"Indice temporale output: {profile_year}\n"
            f"Ore: {N_HOURS}\n\n"
            f"CF wind medio: {mean_wind_cf:.4f}\n"
            f"CF PV medio: {mean_pv_cf:.4f}\n\n"
            f"File salvato in:\n{output_path}"
        ),
    )


# =============================================================================
# ERROR HANDLING
# =============================================================================

if __name__ == "__main__":
    try:
        main()

    except SystemExit:
        raise

    except Exception as exc:
        error_message = (
            f"{type(exc).__name__}: {exc}\n\n"
            "Controlla anche la console di Spyder per i dettagli."
        )

        print("\nERRORE:", flush=True)
        traceback.print_exc()

        try:
            show_error("Errore durante la creazione del file", error_message)
        except Exception:
            pass

        sys.exit(1)
